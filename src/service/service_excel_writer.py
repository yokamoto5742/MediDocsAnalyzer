import os
from datetime import date
from pathlib import Path

import polars as pl
from openpyxl import load_workbook

# 集計結果(analyze_task_data の戻り値)と同じ順に並べる
SHEET_NAMES = ('クラーク業務', 'クラーク以外業務', 'デイリータスク',
               'コミュニケーション', 'コミュニケーション内容', '全項目')


class ExcelResultWriter:
    @staticmethod
    def save_results(analysis_results: tuple[pl.DataFrame, ...], template_path: str, output_dir: str,
                     start_date: date, end_date: date) -> str:
        wb = load_workbook(filename=template_path)

        for sheet_name, result in zip(SHEET_NAMES, analysis_results, strict=True):
            sheet = wb[sheet_name]
            # 1 行目は見出しなので 2 行目から書く
            for row_index, row in enumerate(result.iter_rows(), start=2):
                for column_index, value in enumerate(row, start=1):
                    sheet.cell(row=row_index, column=column_index, value=value)

        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)

        output_filename = f'WILLDOリストまとめ{start_date:%Y%m%d}_{end_date:%Y%m%d}.xlsx'
        output_file_path = output_path / output_filename

        wb.save(output_file_path)

        os.startfile(output_file_path)

        return str(output_file_path)

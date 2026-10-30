from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from service import service_excel_writer
from service.service_database import open_database
from service.service_excel_reader import NoDataError
from service.service_master import Department
from service.service_medical_docs_analyzer import MedicalDocsAnalyzer, split_list
from tests.service.conftest import MakeWorkbook


@pytest.fixture
def startfile(monkeypatch: pytest.MonkeyPatch) -> MagicMock:
    mock = MagicMock()
    monkeypatch.setattr(service_excel_writer.os, 'startfile', mock, raising=False)
    return mock


def database_rows(config_path: Path) -> list[tuple[str, str]]:
    with open_database(str(config_path.parent / 'database.db')) as connection:
        return connection.execute(
            'SELECT received_date, staff_name FROM documents ORDER BY received_date, staff_name').fetchall()


def test_split_list() -> None:
    assert split_list(' 植田, 花﨑 ,,') == ['植田', '花﨑']


def test_import_migrates_legacy_database_then_replaces_source_range(config_path: Path,
                                                                    make_workbook: MakeWorkbook) -> None:
    tmp_path = config_path.parent
    make_workbook(tmp_path / 'legacy.xlsx', [
        ['2025/03/31', 1, '既存', '証明書', '内科', '佐藤', None, '2025/04/01', '沖野'],
        ['2025/04/03', 2, '取消', '診断書', '内科', '佐藤', None, '2025/04/04', '沖野'],
    ])
    make_workbook(tmp_path / 'source.xlsx', [
        [datetime(2025, 4, 1), 3, '山田', '診断書', '内科', '佐藤', None, None, '植田'],
        [datetime(2025, 4, 3), 4, '鈴木', '紹介状', '外科', '田中', None, datetime(2025, 4, 4), '花﨑'],
        [datetime(2025, 4, 5), 5, '田中', '紹介状', '外科', '田中', None, '025/4/6', '花﨑'],
    ])
    analyzer = MedicalDocsAnalyzer()

    plan = analyzer.prepare_import()
    count = analyzer.apply_import(plan)

    assert plan.start_date == date(2025, 4, 1)
    assert plan.legacy is not None
    assert plan.deleted_count == 1
    assert [error.row_number for error in plan.errors] == [4]
    assert count == 2
    assert database_rows(config_path) == [('2025-03-31', '沖野'), ('2025-04-03', '花﨑')]


def test_import_keeps_documents_newer_than_source(config_path: Path, make_workbook: MakeWorkbook) -> None:
    tmp_path = config_path.parent
    make_workbook(tmp_path / 'legacy.xlsx', [
        ['2025/04/01', 1, '既存', '証明書', '内科', '佐藤', None, '2025/04/02', '沖野'],
        ['2025/06/30', 2, '新しい', '証明書', '内科', '佐藤', None, '2025/07/01', '沖野'],
    ])
    make_workbook(tmp_path / 'source.xlsx', [
        ['2025/04/01', 1, '既存', '証明書', '内科', '佐藤', None, '2025/04/02', '植田'],
        ['2025/04/02', 3, '下書き', '証明書', '内科', '佐藤', None, None, None],
    ])
    analyzer = MedicalDocsAnalyzer()

    plan = analyzer.prepare_import()
    analyzer.apply_import(plan)

    assert (plan.start_date, plan.end_date, plan.deleted_count) == (date(2025, 4, 1), date(2025, 4, 2), 1)
    assert database_rows(config_path) == [('2025-04-01', '植田'), ('2025-06-30', '沖野')]


def test_import_skips_legacy_database_when_database_has_documents(config_path: Path,
                                                                  make_workbook: MakeWorkbook) -> None:
    tmp_path = config_path.parent
    make_workbook(tmp_path / 'source.xlsx', [
        ['2025/04/03', 1, '山田', '診断書', '内科', '佐藤', None, '2025/04/04', '植田'],
    ])
    analyzer = MedicalDocsAnalyzer()
    analyzer.apply_import(analyzer.prepare_import())
    make_workbook(tmp_path / 'legacy.xlsx', [
        ['2025/03/31', 1, '既存', '証明書', '内科', '佐藤', None, '2025/04/01', '沖野'],
    ])

    plan = analyzer.prepare_import()

    assert plan.legacy is None
    assert plan.deleted_count == 1
    assert analyzer.apply_import(plan) == 1


def test_prepare_import_rejects_staff_not_in_master(config_path: Path, make_workbook: MakeWorkbook) -> None:
    make_workbook(config_path.parent / 'source.xlsx', [
        ['2025/04/03', 1, '山田', '診断書', '内科', '佐藤', None, '2025/04/04', '未登録'],
    ])

    plan = MedicalDocsAnalyzer().prepare_import()

    assert [error.reason for error in plan.errors] == ['担当者名がマスタに登録されていません']
    assert plan.source.documents == []


def test_prepare_import_raises_when_source_is_empty(config_path: Path, make_workbook: MakeWorkbook) -> None:
    make_workbook(config_path.parent / 'source.xlsx', [])

    with pytest.raises(NoDataError):
        MedicalDocsAnalyzer().prepare_import()


def test_backup_database(config_path: Path) -> None:
    tmp_path = config_path.parent
    with open_database(str(tmp_path / 'database.db')):
        pass

    backup_path = MedicalDocsAnalyzer().backup_database(date(2025, 4, 1))

    assert backup_path == tmp_path / 'backup' / 'database_20250401.db'


def test_run_analysis(config_path: Path, make_workbook: MakeWorkbook, startfile: MagicMock) -> None:
    tmp_path = config_path.parent
    Workbook().save(tmp_path / 'template.xlsx')
    make_workbook(tmp_path / 'source.xlsx', [
        ['2025/03/31', 1, '対象外', '診断書', '内科', '佐藤', None, '2025/04/01', '植田'],
        ['2025/04/01', 2, '山田', '診断書', '内科', '佐藤', None, '2025/04/02', '植田'],
        ['2025/05/10', 3, '鈴木', '紹介状', '外科', '田中', None, '2025/05/11', '花﨑'],
    ])
    analyzer = MedicalDocsAnalyzer()
    analyzer.apply_import(analyzer.prepare_import())

    output_path = analyzer.run_analysis(date(2025, 4, 1), date(2025, 6, 30))

    assert Path(output_path).name == '医療文書作成件数20250401-20250510.xlsx'
    sheet = load_workbook(output_path).active
    assert sheet is not None
    assert [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)] == [
        ['氏名', '合計', '内科', '外科'],
        ['植田', 1, 1, 0],
        ['沖野', 0, 0, 0],
        ['花﨑', 1, 0, 1],
        ['合計', 2, 1, 1],
    ]
    startfile.assert_called_once()


def test_run_analysis_uses_masters(config_path: Path, make_workbook: MakeWorkbook, startfile: MagicMock) -> None:
    tmp_path = config_path.parent
    Workbook().save(tmp_path / 'template.xlsx')
    make_workbook(tmp_path / 'source.xlsx', [
        ['2025/04/01', 1, '山田', '診断書', '眼科', '佐藤', None, '2025/04/02', '植田'],
    ])
    analyzer = MedicalDocsAnalyzer()
    analyzer.apply_import(analyzer.prepare_import())
    ueda, okino, _ = analyzer.list_staff()
    ueda.active = False
    okino.active = False
    analyzer.save_staff(ueda)
    analyzer.save_staff(okino)
    analyzer.save_department(None, Department('眼科', 0))

    output_path = analyzer.run_analysis(date(2025, 4, 1), date(2025, 4, 30))

    sheet = load_workbook(output_path).active
    assert sheet is not None
    assert [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)] == [
        ['氏名', '合計', '眼科', '内科', '外科'],
        ['植田', 1, 1, 0, 0],
        ['花﨑', 0, 0, 0, 0],
        ['合計', 1, 1, 0, 0],
    ]


def test_export_staff(config_path: Path, make_workbook: MakeWorkbook, startfile: MagicMock) -> None:
    tmp_path = config_path.parent
    make_workbook(tmp_path / 'source.xlsx', [
        ['2025/04/01', 1, '山田', '診断書', '内科', '佐藤', None, '2025/04/02', '花﨑'],
    ])
    analyzer = MedicalDocsAnalyzer()
    analyzer.apply_import(analyzer.prepare_import())

    output_path = analyzer.export_staff('花﨑', date(2025, 4, 1), date(2025, 4, 30))

    assert Path(output_path) == tmp_path / 'output' / '医療文書作成件数_花﨑_20250401-20250430.xlsx'
    assert load_workbook(output_path)['個人']['B2'].value == 1


def test_run_analysis_raises_when_no_data_in_period(config_path: Path) -> None:
    with pytest.raises(NoDataError):
        MedicalDocsAnalyzer().run_analysis(date(2025, 4, 1), date(2025, 6, 30))


def test_empty_workbook_without_headers_is_rejected(config_path: Path) -> None:
    Workbook().save(config_path.parent / 'source.xlsx')

    with pytest.raises(ValueError):
        MedicalDocsAnalyzer().prepare_import()

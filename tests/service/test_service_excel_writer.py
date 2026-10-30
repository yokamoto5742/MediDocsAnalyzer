from datetime import date
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from service import service_excel_writer
from service.service_data_analyzer import DocumentCounts, StaffCounts
from service.service_excel_writer import save_results, save_staff_results

TITLE = '医療文書作成件数'
STAFF_NAME = '花﨑'
DEPARTMENT = '内科'
DOCUMENT = '診断書'


def active_sheet(path: Path | str) -> Worksheet:
    sheet = load_workbook(path).active
    assert isinstance(sheet, Worksheet)
    return sheet


def sheet_values(sheet: Worksheet) -> list[list[object]]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_save_results(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    startfile = MagicMock()
    monkeypatch.setattr(service_excel_writer.os, 'startfile', startfile, raising=False)
    template_path = tmp_path / 'template.xlsx'
    Workbook().save(template_path)
    counts = DocumentCounts(
        by_staff_department={('植田', '内科'): 2, ('花﨑', '外科'): 1, ('外部', '内科'): 1, ('植田', '眼科'): 5},
        by_department={'内科': 3, '外科': 1, '眼科': 5},
        total=9,
        first_date=date(2025, 4, 1),
        last_date=date(2025, 5, 20),
    )

    output_path = save_results(counts, ['植田', '花﨑'], ['合計', '内科', '外科'],
                               str(template_path), str(tmp_path / 'output'))

    assert Path(output_path) == tmp_path / 'output' / '医療文書作成件数20250401-20250520.xlsx'
    assert sheet_values(active_sheet(output_path)) == [
        ['医療文書作成件数 2025年04月01日-2025年05月20日', None, None, None],
        ['氏名', '合計', '内科', '外科'],
        ['植田', 2, 2, 0],
        ['花﨑', 1, 0, 1],
        ['合計', 9, 3, 1],
    ]
    startfile.assert_called_once_with(Path(output_path))
    sheet = active_sheet(output_path)
    assert len(getattr(sheet, '_charts')) == 1
    assert sheet.page_setup.orientation == 'landscape'


def test_save_staff_results(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    startfile = MagicMock()
    monkeypatch.setattr(service_excel_writer.os, 'startfile', startfile, raising=False)
    staff_counts = StaffCounts(STAFF_NAME, 3, 1.25, [(DEPARTMENT, 3)], [(DOCUMENT, 3)],
                               [('2025/04', 2), ('2025/05', 1)])

    output_path = save_staff_results(staff_counts, date(2025, 4, 1), date(2025, 5, 31), str(tmp_path / 'output'))

    assert Path(output_path).name == f'{TITLE}_{STAFF_NAME}_20250401-20250531.xlsx'
    sheet = active_sheet(output_path)
    assert sheet['B2'].value == 3
    assert sheet['E2'].value == 1.2
    values = sheet_values(sheet)
    assert values[4][:8] == [DEPARTMENT, 3, None, DOCUMENT, 3, None, '2025/04', 2]
    assert values[5][6:8] == ['2025/05', 1]
    assert len(getattr(sheet, '_charts')) == 1
    startfile.assert_called_once_with(Path(output_path))

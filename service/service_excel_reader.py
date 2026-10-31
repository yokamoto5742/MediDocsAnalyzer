import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

DATE_PATTERN = re.compile(r'\d{4}/\d{1,2}/\d{1,2}')
RECEIVED_DATE_COLUMN = '預り日'
REQUESTED_DATE_COLUMN = '医師依頼日'
STAFF_COLUMN = '担当者名'
PATIENT_ID_COLUMN = '患者ID'
TEXT_COLUMNS = ('文書名', '診療科', '医師名', STAFF_COLUMN)
REQUIRED_COLUMNS = (RECEIVED_DATE_COLUMN, REQUESTED_DATE_COLUMN, *TEXT_COLUMNS, PATIENT_ID_COLUMN)


class NoDataError(Exception):
    """対象となるデータが 1 件もない"""


@dataclass(frozen=True)
class Document:
    received_date: date
    requested_date: date
    document_name: str
    department: str
    doctor_name: str
    staff_name: str
    patient_id: str = ''


@dataclass(frozen=True)
class Draft:
    """医師依頼日か担当者名が空の行(下書き)"""
    received_date: date
    requested_date: date | None
    document_name: str
    department: str
    doctor_name: str
    staff_name: str
    patient_id: str


@dataclass(frozen=True)
class RowError:
    file_name: str
    row_number: int
    patient_id: str
    reason: str
    row_text: str


@dataclass
class SourceData:
    documents: list[Document] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    drafts: list[Draft] = field(default_factory=list)
    first_received_date: date | None = None
    last_received_date: date | None = None


def to_date(value: object) -> date | None:
    """日付・Excel のシリアル値・'YYYY/MM/DD' 形式の文字列を日付にする。空なら None、読めなければ ValueError"""
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return from_excel(value).date()
    text = str(value).strip().split()[0].replace('-', '/')
    if not DATE_PATTERN.fullmatch(text):
        raise ValueError(text)
    return datetime.strptime(text, '%Y/%m/%d').date()


def to_text(value: object) -> str:
    return '' if value is None else str(value).strip()


def row_error_reason(cells: dict[str, object], known_staff: set[str] | None) -> str | None:
    try:
        received = to_date(cells[RECEIVED_DATE_COLUMN])
    except ValueError:
        return '預り日が日付ではありません'
    if received is None:
        return '預り日が空です'
    try:
        requested = to_date(cells[REQUESTED_DATE_COLUMN])
    except ValueError:
        return '医師依頼日が日付ではありません'
    if requested is not None and requested < received:
        return '医師依頼日が預り日より前です'
    staff = to_text(cells[STAFF_COLUMN])
    if requested is not None and staff and known_staff is not None and staff not in known_staff:
        return '担当者名がマスタに登録されていません'
    return None


def read_documents(file_path: str, known_staff: set[str] | None = None) -> SourceData:
    """1 行目を見出しとして読み、医師依頼日と担当者名がある行を文書として返す。

    空欄の行は下書きとして drafts に入れ、日付が読めない行などはエラーとして返す。
    first_received_date / last_received_date には、読めた預り日(下書きを含む)の最も古い日・新しい日を入れる。
    """
    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheet = wb.active
    assert sheet is not None
    rows = sheet.iter_rows(values_only=True)
    headers = [to_text(value) for value in next(rows, ())]
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        wb.close()
        raise ValueError(f'{Path(file_path).name} に列がありません: {", ".join(missing)}')

    column_indexes = {column: headers.index(column) for column in REQUIRED_COLUMNS}
    source = SourceData()
    for row_number, row in enumerate(rows, start=2):
        cells: dict[str, object] = {column: row[index] if index < len(row) else None
                                    for column, index in column_indexes.items()}
        if all(to_text(value) == '' for value in cells.values()):
            continue
        _add_row(source, cells, row_number, Path(file_path).name, known_staff)
    wb.close()
    return source


def _add_row(source: SourceData, cells: dict[str, object], row_number: int, file_name: str,
             known_staff: set[str] | None) -> None:
    # 預り日が読める行は、エラーでも取り込み範囲に含める
    try:
        received = to_date(cells[RECEIVED_DATE_COLUMN])
    except ValueError:
        received = None
    if received is not None:
        source.first_received_date = min(received, source.first_received_date or received)
        source.last_received_date = max(received, source.last_received_date or received)

    reason = row_error_reason(cells, known_staff)
    if reason is not None:
        row_text = ' | '.join(to_text(value) for column, value in cells.items() if column != PATIENT_ID_COLUMN)
        source.errors.append(RowError(file_name, row_number, to_text(cells[PATIENT_ID_COLUMN]), reason, row_text))
        return

    requested = to_date(cells[REQUESTED_DATE_COLUMN])
    staff = to_text(cells[STAFF_COLUMN])
    assert received is not None
    texts = (to_text(cells['文書名']), to_text(cells['診療科']), to_text(cells['医師名']), staff,
             to_text(cells[PATIENT_ID_COLUMN]))
    if requested is None or not staff:
        source.drafts.append(Draft(received, requested, *texts))
        return
    source.documents.append(Document(received, requested, *texts))

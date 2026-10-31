import os
from datetime import date
from pathlib import Path
from typing import cast

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from service.service_data_analyzer import DocumentCounts, StaffCounts

TOTAL_LABEL = '合計'


def write_counts(sheet: Worksheet, counts: DocumentCounts, staff_members: list[str],
                 departments: list[str]) -> None:
    """2 行目に診療科、A列に担当者を並べ、件数を入れる。'合計' 列には行の合計を入れる"""
    sheet['A2'] = '氏名'
    for column_index, department in enumerate(departments, start=2):
        sheet.cell(row=2, column=column_index, value=department)

    total_row = len(staff_members) + 3
    for row_index, staff in enumerate(staff_members, start=3):
        sheet.cell(row=row_index, column=1, value=staff)
    sheet.cell(row=total_row, column=1, value=TOTAL_LABEL)

    for column_index, department in enumerate(departments, start=2):
        if department == TOTAL_LABEL:
            continue
        for row_index, staff in enumerate(staff_members, start=3):
            sheet.cell(row=row_index, column=column_index,
                       value=counts.by_staff_department.get((staff, department), 0))
        sheet.cell(row=total_row, column=column_index, value=counts.by_department.get(department, 0))

    if TOTAL_LABEL in departments:
        total_column = departments.index(TOTAL_LABEL) + 2
        listed_departments = [department for department in departments if department != TOTAL_LABEL]
        for row_index, staff in enumerate(staff_members, start=3):
            staff_total = sum(counts.by_staff_department.get((staff, department), 0)
                              for department in listed_departments)
            sheet.cell(row=row_index, column=total_column, value=staff_total)
        # 一覧にない担当者・診療科の件数も含めた全体の件数
        sheet.cell(row=total_row, column=total_column, value=counts.total)


def save_results(counts: DocumentCounts, staff_members: list[str], departments: list[str],
                 template_path: str, output_dir: str) -> str:
    """集計結果をテンプレートに書いて保存し、保存先のパスを返す"""
    wb = load_workbook(template_path)
    sheet = wb.active
    assert sheet is not None

    sheet['A1'] = (f'医療文書作成件数 {counts.first_date:%Y年%m月%d日}'
                   f'-{counts.last_date:%Y年%m月%d日}')
    write_counts(sheet, counts, staff_members, departments)
    add_overall_chart(sheet, len(staff_members), departments)
    set_print_layout(sheet)

    return save_and_open(wb, output_dir,
                         f'医療文書作成件数{counts.first_date:%Y%m%d}-{counts.last_date:%Y%m%d}.xlsx')


def save_and_open(wb: Workbook, output_dir: str, file_name: str) -> str:
    """保存して出力先フォルダを開き、保存先のパスを返す"""
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    output_file_path = output_path / file_name
    wb.save(output_file_path)

    # フォルダを開く(エクスプローラーは同じフォルダの窓が開いていれば、新しく開かずにそれを前に出す)
    os.startfile(output_path)

    return str(output_file_path)


def set_print_layout(sheet: Worksheet) -> None:
    """A4 横 1 枚に収めて印刷する"""
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1


def make_bar_chart(title: str, horizontal: bool) -> BarChart:
    chart = BarChart()
    chart.type = 'bar' if horizontal else 'col'
    chart.title = title
    chart.width = 24
    chart.height = 10
    # openpyxl 3.1 は軸を既定で隠すため、明示的に表示する
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    return chart


def add_overall_chart(sheet: Worksheet, staff_count: int, departments: list[str]) -> None:
    """担当者ごとの件数を診療科で積み上げた横棒グラフを、表の下に置く"""
    chart = make_bar_chart('担当者別件数', horizontal=True)
    chart.grouping = 'stacked'
    chart.overlap = 100
    # 表と同じく上から担当者を並べる
    chart.x_axis.scaling.orientation = 'maxMin'
    for column_index, department in enumerate(departments, start=2):
        if department != TOTAL_LABEL:
            chart.add_data(Reference(sheet, min_col=column_index, min_row=2, max_row=staff_count + 2),
                           titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=staff_count + 2))
    sheet.add_chart(chart, f'A{staff_count + 5}')


def write_table(sheet: Worksheet, row: int, column: int, headers: tuple[str, str],
                rows: list[tuple[str, int]]) -> None:
    for offset, value in enumerate(headers):
        sheet.cell(row=row, column=column + offset, value=value).font = Font(bold=True)
    for row_offset, (label, count) in enumerate(rows, start=1):
        sheet.cell(row=row + row_offset, column=column, value=label)
        sheet.cell(row=row + row_offset, column=column + 1, value=count)


def save_staff_results(staff_counts: StaffCounts, start_date: date, end_date: date, output_dir: str) -> str:
    """1 人分の集計(診療科別・文書名別・月別と月別グラフ)を新しいブックに書いて保存する"""
    wb = Workbook()
    sheet = cast(Worksheet, wb.active)
    sheet.title = '個人'
    name = staff_counts.staff_name
    sheet['A1'] = f'医療文書作成件数 {name} {start_date:%Y年%m月%d日}-{end_date:%Y年%m月%d日}'
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A2'] = TOTAL_LABEL
    sheet['B2'] = staff_counts.total
    sheet['D2'] = '担当者1人あたりの平均'
    sheet['E2'] = round(staff_counts.average, 1)

    write_table(sheet, 4, 1, ('診療科', '件数'), staff_counts.by_department)
    write_table(sheet, 4, 4, ('文書名', '件数'), staff_counts.by_document)
    write_table(sheet, 4, 7, ('月', '件数'), staff_counts.by_month)
    for column, width in (('A', 14), ('D', 48), ('G', 10)):
        sheet.column_dimensions[column].width = width

    chart = make_bar_chart(f'{name} 月別件数', horizontal=False)
    chart.legend = None
    chart.add_data(Reference(sheet, min_col=8, min_row=4, max_row=4 + len(staff_counts.by_month)),
                   titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=7, min_row=5, max_row=4 + len(staff_counts.by_month)))
    table_rows = max(len(staff_counts.by_department), len(staff_counts.by_document), len(staff_counts.by_month))
    sheet.add_chart(chart, f'A{table_rows + 7}')
    set_print_layout(sheet)

    return save_and_open(wb, output_dir,
                         f'医療文書作成件数_{name}_{start_date:%Y%m%d}-{end_date:%Y%m%d}.xlsx')
